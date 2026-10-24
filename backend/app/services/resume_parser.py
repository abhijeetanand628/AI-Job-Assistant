from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SKILLS = [
    "Python", "JavaScript", "TypeScript", "React", "React Native", "Node.js", "Next.js", "Vue", "Angular", "Svelte",
    "FastAPI", "Django", "Flask", "Express", "Spring", ".NET", "Ruby on Rails", "PHP", "Java", "Kotlin", "Swift", "Go", "Rust", "Scala", "C++", "C#",
    "HTML", "CSS", "Tailwind", "Sass", "GraphQL", "REST", "SQL", "PostgreSQL", "MySQL", "MongoDB", "Redis", "Snowflake", "dbt",
    "AWS", "Azure", "Google Cloud", "Docker", "Kubernetes", "Terraform", "Git", "GitHub Actions", "Linux", "CI/CD",
    "Machine Learning", "Deep Learning", "Natural Language Processing", "TensorFlow", "PyTorch", "scikit-learn", "Pandas", "NumPy", "Apache Spark", "Airflow",
    "Generative AI", "LLM", "LangChain", "LangGraph", "OpenAI", "Gemini", "Anthropic", "AI agents", "Prompt Engineering", "Figma",
    "Data Analysis", "Data Analytics", "Business Intelligence", "Statistics", "Experiment Design", "A/B Testing", "Tableau", "Power BI", "Looker", "Excel",
    "Sales", "Salesforce", "Business Development", "Account Management", "Customer Success", "Customer Support", "Customer Service", "Marketing", "SEO", "Content Strategy", "Public Relations",
    "Product Management", "Project Management", "Agile", "Scrum", "Operations", "Supply Chain", "Finance", "Financial Analysis", "Accounting", "Audit", "Tax", "Risk Management", "Compliance", "ERP", "SAP",
    "Human Resources", "Recruiting", "User Research", "Market Research", "Strategy", "Stakeholder Management", "Negotiation", "Communication", "Research", "Instructional Design",
]
_SKILL_PATTERN = re.compile(r"(?<![\w])(?:" + "|".join(re.escape(skill) for skill in sorted(SKILLS, key=len, reverse=True)) + r")(?![\w])", re.I)


def find_skills(text: str) -> list[str]:
    found = {match.group(0).casefold() for match in _SKILL_PATTERN.finditer(text)}
    return [skill for skill in SKILLS if skill.casefold() in found]


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        elif suffix == ".docx":
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        else:
            raise ValueError("Upload a PDF or DOCX file.")
    except Exception as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError("Could not read this resume. Check that the file is a valid, text-based PDF or DOCX.") from exc
    text = re.sub(r"[\t\r]+", " ", text)
    if len(text.strip()) < 30:
        raise ValueError("No readable text found. Scanned PDF resumes are not supported yet.")
    return text[:100_000]


def basic_profile(text: str) -> dict:
    lines = [re.sub(r"\s+", " ", line).strip(" •|-") for line in text.splitlines() if line.strip()]
    email = next(iter(re.findall(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text)), "")
    phone = next(iter(re.findall(r"(?:\+?\d[\d .()/-]{7,}\d)", text)), "")
    name = next((line for line in lines[:8] if len(line.split()) in (2, 3, 4) and not re.search(r"@|resume|curriculum|linkedin|github|\d", line, re.I)), "")
    low = text.lower()
    skills = find_skills(text)
    exp_match = re.search(r"(\d+(?:\.\d+)?)\+?\s+years?\s+(?:of\s+)?experience", low)
    experience = float(exp_match.group(1)) if exp_match else _experience_from_dates(lines)
    section = lambda words: _section(lines, words)
    designation = _labeled_value(text, ("designation", "current role", "current title", "professional title", "job title", "position"))
    return {"name": name, "email": email, "phone": phone, "designation": designation, "title": "", "location": "", "skills": skills, "experience": experience, "education": section(["education", "academic background", "education and qualifications"]), "projects": section(["projects", "selected work", "personal projects"])}


def _labeled_value(text: str, labels: tuple[str, ...]) -> str:
    alternatives = "|".join(re.escape(label) for label in labels)
    match = re.search(rf"(?im)^\s*(?:{alternatives})\s*[:\-–]\s*([^\n|]+)", text)
    return re.sub(r"\s+", " ", match.group(1)).strip(" ,.;")[:120] if match else ""


def _experience_from_dates(lines: list[str]) -> float:
    """Estimate experience from date ranges within a work history section."""
    start = next((i for i, line in enumerate(lines) if re.search(r"\b(experience|employment|work history)\b", line, re.I)), None)
    if start is None:
        return 0
    work_lines = []
    for line in lines[start + 1 :]:
        if re.fullmatch(r"[A-Z][A-Z &/]{2,}", line) and not re.search(r"EXPERIENCE|EMPLOYMENT|WORK HISTORY", line):
            break
        work_lines.append(line)
    month_names = {name.lower(): number for number, name in enumerate(("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"), 1)}
    month_pattern = r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
    ranges = re.findall(rf"\b({month_pattern})?\s*(19\d{{2}}|20\d{{2}})\s*(?:-|–|—|to)\s*({month_pattern})?\s*(present|current|19\d{{2}}|20\d{{2}})\b", " ".join(work_lines), re.I)
    now = datetime.now()
    periods = []
    for start_month, start_year, end_month, end_year in ranges:
        start_month_num = month_names.get(start_month[:3].lower(), 1) if start_month else 1
        if end_year.lower() in {"present", "current"}:
            end_year_num, end_month_num = now.year, now.month
        else:
            end_year_num = int(end_year)
            end_month_num = month_names.get(end_month[:3].lower(), 1) if end_month else 1
        start_idx = int(start_year) * 12 + start_month_num
        end_idx = end_year_num * 12 + end_month_num
        if end_idx > start_idx:
            periods.append((start_idx, end_idx))
    if not periods:
        return 0
    merged = []
    for start, end in sorted(periods):
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((start, end))
    total_months = sum(end - start for start, end in merged)
    return round(total_months / 12, 1)


def _section(lines: list[str], names: list[str]) -> list[str]:
    def heading_name(line: str) -> tuple[str, str]:
        cleaned = line.strip().lstrip("•*- ").strip()
        for name in names:
            match = re.match(rf"^{re.escape(name)}\b\s*(?::|\||-|–)?\s*(.*)$", cleaned, re.I)
            if match:
                return name, match.group(1).strip()
        return "", ""

    found = next(((i, heading_name(line)[1]) for i, line in enumerate(lines) if heading_name(line)[0]), None)
    start = found[0] if found else None
    if start is None:
        return []
    result = []
    if found[1]:
        result.append(found[1])
    boundaries = ("experience", "employment", "work history", "education", "academic background", "projects", "selected work", "skills", "certifications", "summary", "objective", "publications", "awards")
    for line in lines[start + 1 : start + 13]:
        cleaned = line.strip().lstrip("•*- ").strip()
        lowered = cleaned.lower().rstrip(":")
        if any(lowered == boundary or lowered.startswith(boundary + ":") for boundary in boundaries):
            break
        if cleaned:
            result.append(cleaned[:240])
    return result
